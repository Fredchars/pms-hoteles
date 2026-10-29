"""Genera la guía de pruebas por escenarios (Word + PDF) y la lista de chequeo en Excel (.xlsx)."""
import sys
sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from docs_comun import *  # noqa
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

U = {"rec": "recepcion1.col@pms-demo.test", "rec2": "recepcion2.col@pms-demo.test", "ger1": "gerencia1.col@pms-demo.test",
     "ger2": "gerencia2.col@pms-demo.test", "hk1": "housekeeping1.col@pms-demo.test", "hk2": "housekeeping2.col@pms-demo.test",
     "adm": "admin@pms-demo.test"}
# (id, grupo, título, usuario, pasos, esperado)
E = [
 ("E01", "Acceso", "Ingreso correcto y contraseña mala", U["rec"],
  ["Abre https://pms-hoteles-atm.streamlit.app/.", "Escribe un correo válido con una contraseña equivocada y pulsa Entrar.", "Ingresa con la contraseña temporal correcta (si ya creaste la tuya, usa la tuya; el primer ingreso se prueba en E33)."],
  "Con la contraseña mala sale un mensaje claro y no entra. Con la correcta entra (si es la temporal, primero pide crear la contraseña nueva; ya creada, ve el menú de recepción)."),
 ("E02", "Acceso", "Bloqueo por intentos fallidos", U["hk2"],
  ["Falla la contraseña 5 veces seguidas.", "Intenta entrar con la contraseña correcta."],
  "Después del 5.º fallo avisa que el acceso está bloqueado unos minutos, aunque la contraseña sea correcta. Queda en Auditoría (ingresar como gerencia para verlo)."),
 ("E03", "Acceso", "La sesión sobrevive a recargar la página", U["rec"],
  ["Ingresa.", "Pulsa F5 (recargar).", "Cambia de pantalla varias veces y recarga otra vez."],
  "No pide ingresar de nuevo. Sigue en la misma pantalla y en el mismo hotel. (Si aún no has creado tu contraseña, recargar te mantiene en «Crea tu contraseña nueva»: ver E37.)"),
 ("E04", "Acceso", "Sin acceso de demostración abierto", "—",
  ["Mira la pantalla de ingreso de la dirección de producción."],
  "No aparecen botones de «entrar sin contraseña». Solo correo y contraseña."),
 ("E05", "Roles", "Cada rol ve solo lo suyo", f"{U['rec']}, {U['ger1']}, {U['hk1']}, {U['adm']}",
  ["Ingresa con cada usuario y revisa el menú de la izquierda.", "Compara con la tabla de permisos del instructivo."],
  "Recepción ve el Plano de habitaciones pero no Reportes, Configuración ni Auditoría. Housekeeping solo ve Housekeeping y Bloqueos (ni plano ni precios). Gerencia ve Reportes. Administrador ve Configuración."),
 ("E06", "Roles", "Un hotel no ve datos de otro", U["rec"],
  ["Ingresa con recepción de un hotel.", "Revisa el selector de hotel."],
  "Solo aparece su hotel. Las reservas y huéspedes son solo de ese hotel."),
 ("E07", "Reservas", "Crear una reserva normal", U["rec"],
  ["Reservas → Nueva reserva.", "Elige un huésped con consentimiento, tipo Doble matrimonial, 3 noches desde la semana próxima.", "Revisa el desglose y guarda."],
  "Se crea con código, habitación asignada, 3 noches con precio, impuestos del hotel y total. Aparece en la Lista y en el Calendario."),
 ("E08", "Reservas", "No se puede vender dos veces la misma habitación", U["rec"],
  ["Reserva todas las habitaciones de un tipo para unas fechas (hay 9 Dobles matrimoniales por hotel: 3 por planta).", "Intenta una más para las mismas fechas."],
  "Al agotarse el tipo, el sistema avisa «no hay disponibilidad» y no guarda. Nunca hay dos reservas en la misma habitación y noche."),
 ("E09", "Reservas", "Precio congelado por noche", f"{U['rec']} y {U['ger1']}",
  ["Con recepción, crea una reserva y anota el precio por noche.", "Con gerencia, cambia la tarifa de ese tipo (Tarifas).", "Vuelve a la reserva (Gestionar → Noches y precio congelado).", "Crea otra reserva igual."],
  "La primera reserva conserva su precio. La segunda usa el precio nuevo."),
 ("E10", "Reservas", "Modificar fechas de una reserva", U["rec"],
  ["Gestionar una reserva → Cambiar fechas → agrega una noche.", "Intenta mover a fechas sin disponibilidad."],
  "Las noches antiguas conservan precio y la nueva usa la tarifa vigente. Si no hay habitación, avisa y no cambia nada."),
 ("E11", "Reservas", "Falta el consentimiento (Ley 1581)", U["rec"],
  ["Crea un huésped nuevo sin registrar consentimiento.", "Intenta reservarle."],
  "Avisa que falta el consentimiento. Con el consentimiento registrado, deja continuar."),
 ("E12", "Impuestos", "Impuestos distintos por hotel", "recepcion1.col@pms-demo.test, recepcion1.rd@pms-demo.test, recepcion1.pe@pms-demo.test",
  ["Con la recepción de cada uno de los 3 hoteles, crea una reserva igual (mismo tipo y 2 noches).", "Compara las líneas de impuestos y la moneda."],
  "Colombia en COP con su impuesto de ejemplo; República Dominicana en DOP con dos cargos; el tercer país en PEN. Cada total usa el % de su hotel."),
 ("E13", "Estancia", "Check-in y check-out", U["rec"],
  ["Toma una reserva confirmada para hoy y pulsa Hacer check-in.", "Intenta check-in de una reserva de otra fecha.", "Con saldo pendiente, pulsa Hacer check-out.", "Registra el pago y vuelve a pulsar check-out."],
  "El check-in de otra fecha se rechaza. Con saldo pendiente el check-out se rechaza. Pagado, el check-out funciona y la habitación queda Sucia."),
 ("E14", "Estancia", "Salida forzada con saldo (solo gerencia)", U["ger1"],
  ["Con una reserva alojada con saldo, ingresa como gerencia.", "Pulsa Forzar salida con saldo."],
  "La salida se hace y queda registrada en Auditoría. Recepción no tiene ese botón."),
 ("E15", "Estancia", "Marcar no-show", U["rec"],
  ["Toma una reserva confirmada cuya entrada ya pasó.", "Pulsa Marcar no-show."],
  "Pasa a «No se presentó» y la habitación queda libre en el calendario."),
 ("E16", "Aprobaciones", "Cancelar requiere gerencia", f"{U['rec']}, {U['ger1']}, {U['ger2']}",
  ["Recepción solicita la cancelación con un motivo.", "Gerencia 1 revisa en Cancelaciones pendientes y aprueba.", "Repite: Gerencia 1 solicita una cancelación e intenta aprobarla él mismo.", "Aprueba con Gerencia 2."],
  "Recepción no puede cancelar sola. Quien pide la cancelación no puede aprobarla; otra persona de gerencia sí. Todo queda en Auditoría."),
 ("E17", "Aprobaciones", "Descuento requiere gerencia", f"{U['rec']}, {U['ger1']}",
  ["Recepción solicita un descuento del 10 % con motivo.", "Mira el total (no debe cambiar).", "Gerencia aprueba en Descuentos pendientes."],
  "El descuento pendiente no cambia el total. Al aprobar, el total baja. Un descuento mayor que los cargos se rechaza."),
 ("E18", "Facturación", "Extras, pagos y anulaciones", U["rec"],
  ["Agrega un extra con cantidad y precio.", "Registra un pago parcial en efectivo.", "Anula el pago con motivo."],
  "El saldo cambia en cada paso. Las líneas anuladas quedan visibles como anuladas, no se borran."),
 ("E19", "Facturación", "Comprobante PDF", U["rec"],
  ["Abre la cuenta de una reserva con cargos y pagos.", "Descarga el comprobante en PDF."],
  "El PDF muestra cargos, impuestos, pagos y saldo en la moneda del hotel, y dice claramente que NO es una factura legal."),
 ("E20", "Housekeeping", "Flujo de limpieza", U["hk1"],
  ["Entra a Housekeeping y filtra las Sucias.", "Pasa una a En limpieza y luego a En inspección.", "Intenta darla por Lista con el mismo usuario.", "Hazlo con Housekeeping 2."],
  "El mismo usuario no puede dar por lista la habitación que pasó a inspección. Otro usuario sí."),
 ("E21", "Housekeeping", "Housekeeping no ve precios ni pagos", U["hk1"],
  ["Revisa todas las pantallas del menú de housekeeping."],
  "No hay precios, pagos ni datos personales de huéspedes en ninguna pantalla."),
 ("E22", "Housekeeping", "Fuera de servicio pide nota y check-in exige habitación lista", f"{U['hk1']} y {U['rec']}",
  ["Pasa una habitación a Fuera de servicio sin escribir nota.", "Repite con nota.", "Intenta hacer check-in en una habitación que no esté Lista."],
  "Sin nota no deja. Con nota queda fuera de servicio. El check-in en habitación no lista se rechaza."),
 ("E23", "Bloqueos", "Bloquear una habitación", U["ger1"],
  ["Bloqueos → Nuevo bloqueo con motivo y fechas.", "Intenta reservar ese tipo en esas fechas.", "Intenta bloquear una habitación con noches ya vendidas."],
  "La habitación bloqueada no se vende. Bloquear noches vendidas se rechaza con un mensaje que cita la reserva."),
 ("E24", "Calendario", "Calendario de solo lectura", U["rec"],
  ["Abre Calendario.", "Intenta arrastrar o hacer clic en una barra de reserva."],
  "Muestra habitaciones por días con colores por estado. No permite cambiar nada; los cambios se hacen en Reservas."),
 ("E25", "Resultados", "Dashboard del día", U["rec"],
  ["Abre el Dashboard del día.", "Haz clic en los enlaces de una llegada y de una habitación sucia."],
  "Muestra llegadas, salidas, ocupación y limpieza del día. Los enlaces llevan directo a la reserva o a Housekeeping."),
 ("E26", "Resultados", "Reportes de ocupación e ingresos", f"{U['ger1']} y {U['adm']}",
  ["Abre Reportes con un rango de fechas.", "Descarga el CSV.", "Repite con admin@pms-demo.test y cambia de hotel en el menú."],
  "Gráficos y tabla por mes en la moneda del hotel. El CSV abre en Excel. Al cambiar de hotel cambian cifras y moneda."),
 ("E27", "Datos", "Huésped: consentimiento y anonimizar", U["rec"],
  ["Registra y luego revoca un consentimiento.", "Anonimiza un huésped de prueba marcando la casilla de confirmación."],
  "Se registra cada cambio. Tras anonimizar, el nombre y los datos personales desaparecen y las reservas se conservan."),
 ("E28", "Administración", "Configurar hotel, impuestos y habitaciones", U["adm"],
  ["Crea un hotel nuevo con otra moneda.", "Agrega un impuesto con nombre y %.", "Crea un tipo de habitación y una habitación con camas."],
  "El hotel nuevo aparece en el selector, con su moneda. La habitación se puede reservar."),
 ("E29", "Administración", "Usuarios", U["adm"],
  ["Crea un usuario de recepción.", "Desactívalo e intenta ingresar con él.", "Cambia su contraseña."],
  "El usuario desactivado no entra. La contraseña nueva funciona, y esa persona deberá crear la suya al ingresar. Nunca se muestra ninguna contraseña guardada."),
 ("E30", "Seguridad", "Auditoría", U["ger1"],
  ["Haz varias acciones (reserva, cancelación, descuento).", "Abre Auditoría."],
  "Cada acción aparece con quién, cuándo y qué. No hay botones para editar ni borrar."),
 ("E31", "Avisos", "Avisos de datos ficticios y simulaciones", U["rec"],
  ["Recorre todas las pantallas.", "Usa «Simular envío de confirmación» y registra un pago."],
  "Siempre hay un aviso visible de datos ficticios. Correos, pagos e integraciones dicen claramente que son simulados. Abajo aparece «Prototipo diseñado para ATm Consulting» con el logo."),
 ("E32", "Aspecto", "Fondo blanco y uso en pantalla pequeña", U["rec"],
  ["Mira varias pantallas en computador.", "Repite en el celular o con la ventana angosta."],
  "Fondo blanco y limpio, textos legibles, sin elementos cortados."),
 ("E33", "Primer ingreso", "Primer ingreso correcto: crear la contraseña propia", U["rec2"],
  ["Ingresa con el correo y la contraseña temporal que te entregaron y pulsa Entrar.", "Comprueba que solo aparece «Crea tu contraseña nueva» y que no hay menú de pantallas.",
   "Escribe una contraseña nueva de 10 o más caracteres, distinta de la temporal, y repítela en la confirmación.", "Pulsa «Guardar y entrar».", "Pulsa «Cerrar sesión» e ingresa con la contraseña nueva. Intenta ingresar con la temporal."],
  "Tras guardar, entra al sistema y ve el menú de su rol. Con la contraseña nueva entra directo, sin volver a pedir crear otra. La temporal ya no funciona. En Auditoría (gerencia) aparece el cambio de contraseña, sin mostrar contraseñas."),
 ("E34", "Primer ingreso", "Contraseña corta o confirmación distinta: rechazada", U["hk2"],
  ["Ingresa con la contraseña temporal para llegar a «Crea tu contraseña nueva».", "Escribe una contraseña de menos de 10 caracteres (por ejemplo 6) en los dos campos y pulsa «Guardar y entrar».",
   "Escribe una contraseña válida en el primer campo y otra distinta en la confirmación; pulsa «Guardar y entrar»."],
  "En ambos casos aparece un aviso rojo claro («al menos 10 caracteres» / «la confirmación no coincide»), no se guarda nada y se sigue viendo solo la pantalla de contraseña nueva."),
 ("E35", "Primer ingreso", "Contraseña igual a la temporal: rechazada", U["ger2"],
  ["Ingresa con la contraseña temporal.", "En «Crea tu contraseña nueva» escribe la misma contraseña temporal en los dos campos (debe tener 10 o más caracteres) y pulsa «Guardar y entrar»."],
  "Aparece el aviso «La contraseña nueva no puede ser igual a la temporal», no se guarda nada y la persona sigue en la misma pantalla."),
 ("E36", "Primer ingreso", "Otras pantallas bloqueadas antes de crear la contraseña", U["hk1"],
  ["Ingresa con la contraseña temporal (sin crear aún la nueva).", "Mira el lado izquierdo: no debe haber menú de pantallas.",
   "Intenta abrir a mano las direcciones de otras pantallas (por ejemplo, la dirección de la aplicación terminada en /housekeeping o /reservas)."],
  "Solo se ve «Crea tu contraseña nueva» (y el botón Cerrar sesión). Ninguna otra pantalla, dato ni precio es accesible hasta guardar la contraseña nueva."),
 ("E37", "Primer ingreso", "Recargar antes de crear la contraseña", U["rec"],
  ["Ingresa con la contraseña temporal y llega a «Crea tu contraseña nueva».", "Pulsa F5 (recargar).", "Crea la contraseña nueva y pulsa «Guardar y entrar», luego recarga otra vez."],
  "Tras la primera recarga sigue en «Crea tu contraseña nueva» sin pedir ingresar de nuevo y sin dar acceso al sistema. Después de guardar, la recarga mantiene la sesión normal."),
 ("E38", "Primer ingreso", "Administrador genera una contraseña temporal y la sesión anterior se cierra", f"{U['adm']} y {U['rec2']}",
  ["Con recepción 2 ya con contraseña propia, deja su sesión abierta en un navegador.", "Con el administrador: Configuración → Usuarios → Generar contraseña temporal junto a recepcion2.col. Copia la contraseña que aparece.",
   "Pulsa otro botón o recarga la página del administrador.", "En el navegador de recepción 2, pulsa cualquier opción o recarga.", "Ingresa como recepción 2 con la temporal y crea una contraseña nueva."],
  "La contraseña temporal se muestra una sola vez y desaparece al recargar o pulsar otro botón. La tabla de usuarios marca «Sí» en «Debe crear su contraseña». La sesión anterior de recepción 2 se cierra y vuelve a la pantalla de ingreso. Con la temporal pide crear la contraseña nueva. En Auditoría queda quién la generó y para quién, sin la contraseña. El administrador no puede generarla para sí mismo."),

 ("E39", "Plano de habitaciones", "El plano se ve con hotel, planta y leyenda", U["rec"],
  ["Ingresa y abre «Plano de habitaciones» en el menú.", "Mira los selectores de hotel y de planta y cambia entre Planta 1, 2 y 3.", "Revisa la leyenda y los números de arriba (disponibles, ocupadas, en opción, no operativas, capacidad insuficiente)."],
  "Cada planta muestra sus 6 habitaciones dibujadas como tarjetas, con número, tipo, capacidad, metros y precio por noche. La leyenda explica los siete colores (disponible, ocupada, en opción, bloqueada, fuera de servicio, seleccionada, capacidad insuficiente). Las de vista al mar quedan junto a la franja «MAR». Recepción solo ve su hotel."),
 ("E40", "Plano de habitaciones", "Los colores cambian con las fechas", f"{U['rec']} y {U['ger1']}",
  ["Con recepción, crea una reserva confirmada de 2 noches desde mañana en una habitación concreta (en «Reservas», casilla «Habitación (opcional)»).", "En el plano pon esas mismas fechas de entrada y salida.", "Pon como entrada el día de salida de esa reserva.",
   "Crea otra reserva en estado «opción» y un bloqueo (gerencia, pantalla «Bloqueos») en otras dos habitaciones y vuelve al plano con esas fechas."],
  "Con las mismas fechas, la habitación reservada sale roja «Ocupada». Con entrada el día de salida vuelve a verse verde (la noche de salida está libre). La de opción sale amarilla «En opción» y la bloqueada rayada «Bloqueada». El resumen de arriba cuenta cada estado."),
 ("E41", "Plano de habitaciones", "Fuera de servicio y capacidad insuficiente", f"{U['hk1']} y {U['rec']}",
  ["Con housekeeping, pasa una habitación a «Fuera de servicio» (con nota).", "Con recepción, abre el plano con 2 adultos y mira la habitación Individual (capacidad 1) y la que pusiste fuera de servicio.", "Sube a 5 adultos y mira las demás habitaciones."],
  "La habitación fuera de servicio sale con rayado gris y sin botón «Elegir». La Individual sale en gris claro «Capacidad insuficiente» con 2 adultos, y las demás de menor capacidad también al subir las personas. Ninguna de ellas se puede elegir."),
 ("E42", "Plano de habitaciones", "Filtros: tipo, vista al mar y ocultar no disponibles", U["rec"],
  ["Elige un tipo de habitación en «Tipo de habitación».", "Vuelve a «Todos los tipos» y activa «🌊 Solo vista al mar».", "Activa «Ocultar no disponibles» con un día que tenga habitaciones ocupadas.", "Combina los tres filtros y cambia de planta."],
  "Cada filtro deja en pantalla solo lo que corresponde: por tipo, solo ese tipo; con vista al mar, solo las habitaciones marcadas con 🌊 (en el piloto, x01 y x02 de cada planta); con «Ocultar no disponibles», solo las verdes. Si no queda ninguna, un aviso lo dice."),
 ("E43", "Plano de habitaciones", "Elegir una habitación y ver el detalle con el desglose del precio", U["rec"],
  ["Pulsa «Elegir» en una habitación verde.", "Lee el panel de detalle a la derecha.", "Compara el total con el de «Reservas» para el mismo tipo y fechas.", "Cambia las fechas o el número de personas y observa el panel."],
  "La tarjeta se pinta azul «Seleccionada». El panel muestra planta, metros, capacidad, vista, camas, cada noche con su precio, subtotal, impuestos del hotel (con su %), total y el aproximado por persona y noche. El total coincide con el que calcula «Reservas». Si la habitación deja de estar disponible por el cambio, la selección se quita sola."),
 ("E44", "Plano de habitaciones", "Continuar con la reserva: datos prellenados", U["rec"],
  ["En el plano elige una habitación verde con fechas y personas distintas de las que trae por defecto.", "Pulsa «Continuar con la reserva →».", "En «Reservas» revisa tipo, habitación, fechas, adultos y niños.", "Elige un huésped con consentimiento y pulsa «Guardar reserva».", "Vuelve a «Plano de habitaciones»."],
  "Se abre «Reservas» en «Nueva reserva» con un aviso azul «Datos traídos del Plano de habitaciones» y todo ya escrito. No se pierde la sesión ni hay que ingresar de nuevo. Al guardar, la reserva queda en esa habitación exacta y, al volver al plano, aparece roja «Ocupada» para esas fechas."),
 ("E45", "Plano de habitaciones", "Nadie la reservó entretanto: no se vende dos veces", f"{U['rec']} y {U['rec2']}",
  ["Con recepción 1 abre el plano y elige una habitación verde, sin pulsar aún «Continuar».", "Con recepción 2 (otro navegador) crea una reserva en esa misma habitación y fechas.", "Con recepción 1 pulsa «Continuar con la reserva →»."],
  "El sistema avisa que la habitación ya no está disponible y no abre «Reservas» con datos viejos. Si de todos modos se intenta guardar una reserva en una habitación ya vendida esa noche, «Reservas» la rechaza: nunca hay dos reservas en la misma habitación y noche."),
 ("E46", "Plano de habitaciones", "Reservas: habitación opcional y asignación automática", U["rec"],
  ["En «Reservas» → «Nueva reserva» deja «Habitación (opcional)» en «Asignar automáticamente» y guarda una reserva.", "Crea otra y elige una habitación concreta de la lista.", "Intenta elegir una habitación y luego cambia el tipo de habitación."],
  "Con asignación automática el sistema elige una habitación libre (como siempre). Con una habitación elegida, la reserva queda en ella. Al cambiar el tipo, la lista de habitaciones se actualiza y no queda una habitación de otro tipo seleccionada."),
 ("E47", "Plano de habitaciones", "Sin tarifa, cierre de ventas o estancia mínima", f"{U['ger1']} y {U['rec']}",
  ["Con gerencia, crea una tarifa con «cierre de ventas» o con estancia mínima de 3 noches para un tipo y unas fechas.", "Con recepción, abre el plano con esas fechas (1 o 2 noches) y elige una habitación de ese tipo.", "Prueba también unas fechas lejanas sin tarifa."],
  "El panel muestra un aviso claro (ventas cerradas, estancia mínima o sin tarifa) en lugar del precio y el botón «Continuar con la reserva →» queda desactivado. Con fechas válidas se activa."),
 ("E48", "Plano de habitaciones", "Administrador y gerencia: permisos", f"{U['adm']}, {U['ger1']}, {U['hk1']}",
  ["Con el administrador abre el plano, cambia de hotel y elige una habitación verde.", "Con gerencia abre el plano y continúa con una reserva.", "Con housekeeping mira el menú."],
  "El administrador ve todos los hoteles y el detalle y precio, pero NO tiene el botón «Continuar con la reserva» (solo consulta). Gerencia lo usa igual que recepción. Housekeeping no tiene «Plano de habitaciones» en el menú. Al cambiar de hotel en el plano, el selector de la barra lateral cambia con él."),
 ("E49", "Plano de habitaciones", "Vista al mar configurable y sin acceso antes de crear la contraseña", f"{U['adm']} y {U['rec2']}",
  ["Con el administrador: Configuración → Habitaciones y camas: crea una habitación con la casilla «Vista al mar» marcada y revisa la columna «Al mar» de la tabla.", "Abre el plano con «Solo vista al mar».",
   "Con un usuario que aún tenga contraseña temporal (E36), ingresa e intenta abrir la dirección terminada en /plano."],
  "La habitación nueva aparece con 🌊 y en el filtro de vista al mar. Antes de crear la contraseña nueva solo se ve «Crea tu contraseña nueva»: el plano no abre."),
]


# ---------- Word ----------
d = nuevo_documento("Guía de pruebas por escenarios", "PMS hotelero (piloto) · qué probar, cómo y qué debe pasar")
parrafo(d, "Versión 1.2 · 28 de octubre de 2026 · Preparado para ATm Consulting, Colombia", color=GRIS, tam=10)
nota(d, "Todo es ficticio y las confirmaciones, los pagos y las integraciones están simuladas. Marca los resultados en el archivo Excel «checklist-pruebas-pms.xlsx».", "Importante")
salto(d)
d.add_heading("Cómo probar", 1)
lista(d, ["La aplicación está en https://pms-hoteles-atm.streamlit.app/. Pide a ATm Consulting las contraseñas TEMPORALES de los usuarios de prueba (se entregan aparte; no están escritas en este documento). Cada persona crea su propia contraseña (mínimo 10 caracteres) la primera vez que ingresa: se prueba en los escenarios E33 a E38. Los escenarios E34 a E37 necesitan un usuario que todavía tenga su contraseña temporal; si ya creó la suya, pide al administrador que le genere otra temporal (escenario E38).",
          "Hay 3 hoteles ficticios: Colombia (COP), República Dominicana (DOP) y otro país (PEN). Cada uno tiene 18 habitaciones y usuarios de recepción, gerencia y housekeeping.",
          "Los correos de los usuarios tienen esta forma: recepcion1.col@pms-demo.test, gerencia2.rd@pms-demo.test, housekeeping1.pe@pms-demo.test. El administrador es admin@pms-demo.test.",
          "Haz los escenarios en orden. Para cada uno, marca en el Excel «Pasó», «Falló» o «No aplica» y escribe lo que viste si falló.",
          "Si algo falla, anota: escenario, usuario, hotel, qué hiciste, qué esperabas y qué pasó (si puedes, una captura de pantalla)."], numerada=True)
d.add_heading("Usuarios de prueba", 2)
tabla(d, ["Rol", "Correo de ejemplo (hotel Colombia)", "Para qué sirve"], [
    ["Recepción", "recepcion1.col@pms-demo.test / recepcion2.col@pms-demo.test", "Reservas, check-in/out, cobros, huéspedes"],
    ["Gerencia", "gerencia1.col@pms-demo.test / gerencia2.col@pms-demo.test", "Aprobar cancelaciones y descuentos, tarifas, reportes (dos personas para probar la separación de funciones)"],
    ["Housekeeping", "housekeeping1.col@pms-demo.test / housekeeping2.col@pms-demo.test", "Tablero de limpieza"],
    ["Administrador", "admin@pms-demo.test", "Configuración de hoteles, impuestos, habitaciones y usuarios"]], anchos=[3, 7.5, 6])
d.add_heading("Mapa de escenarios", 2)
tabla(d, ["Grupo", "Escenarios"], [[g, ", ".join(e[0] for e in E if e[1] == g)] for g in dict.fromkeys(e[1] for e in E)], anchos=[4, 12])
imagen(d, diagrama("flujo_prueba", ["Probar\nel escenario", "¿Pasó lo\nesperado?", "Marcar en\nel Excel", "Si falló:\nanotar y\nreportar"], colores={2: "#e6f4ea", 3: "#fff3e0"}, cols=4, alto=1.3, ancho=10), 14)
salto(d)
grupo = None
for (i, g, t, u, pasos, esp) in E:
    if g != grupo:
        d.add_heading(lim(g), 1); grupo = g
    d.add_heading(f"{i} · {lim(t)}", 3)
    parrafo(d, f"Usuario: {u}", cursiva=True, color=GRIS, tam=10)
    parrafo(d, "Pasos:", negrita=True)
    lista(d, pasos, numerada=True)
    nota(d, esp, "Debe pasar")
ruta, pdf = guardar_y_pdf(d, "guia-de-pruebas-pms")
print(ruta, pdf)

# ---------- Excel ----------
wb = Workbook(); ws = wb.active; ws.title = "Checklist"
cab = ["ID", "Grupo", "Escenario", "Usuario(s)", "Pasos", "Resultado esperado", "Resultado", "Observaciones", "Fecha", "Responsable"]
ws.append(["PMS hotelero (piloto) · Checklist de pruebas · Prototipo diseñado para ATm Consulting · datos 100 % ficticios"])
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(cab)); ws["A1"].font = Font(bold=True, size=13, color="1F3A5F")
ws.append([]); ws.append(cab)
borde = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
for c in range(1, len(cab) + 1):
    x = ws.cell(row=3, column=c); x.font = Font(bold=True); x.fill = PatternFill("solid", fgColor="E8F0FE"); x.border = borde
    x.alignment = Alignment(wrap_text=True, vertical="center")
for (i, g, t, u, pasos, esp) in E:
    ws.append([i, lim(g), lim(t), u, "\n".join(f"{n}. {lim(p)}" for n, p in enumerate(pasos, 1)), esp, "Pendiente", "", "", ""])
for r in range(4, 4 + len(E)):
    for c in range(1, len(cab) + 1):
        x = ws.cell(row=r, column=c); x.alignment = Alignment(wrap_text=True, vertical="top"); x.border = borde
for c, a in enumerate([7, 14, 30, 30, 55, 50, 12, 35, 12, 18], 1):
    ws.column_dimensions[get_column_letter(c)].width = a
dv = DataValidation(type="list", formula1='"Pendiente,Pasó,Falló,No aplica"', allow_blank=False); ws.add_data_validation(dv)
dv.add(f"G4:G{3 + len(E)}")
from openpyxl.formatting.rule import CellIsRule
rango = f"G4:G{3 + len(E)}"
ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"Pasó"'], fill=PatternFill("solid", bgColor="C8E6C9")))
ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"Falló"'], fill=PatternFill("solid", bgColor="FFCDD2")))
ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"Pendiente"'], fill=PatternFill("solid", bgColor="FFF9C4")))
ws.freeze_panes = "D4"; ws.auto_filter.ref = f"A3:{get_column_letter(len(cab))}{3 + len(E)}"
res = wb.create_sheet("Resumen")
res.append(["Resumen automático"]); res["A1"].font = Font(bold=True, size=13, color="1F3A5F"); res.append([])
n = 3 + len(E)
for k in ("Pendiente", "Pasó", "Falló", "No aplica"):
    res.append([k, f'=COUNTIF(Checklist!G4:G{n},"{k}")'])
res.append(["Total", f"=COUNTA(Checklist!A4:A{n})"]); res.append(["% que pasó", f'=IFERROR(COUNTIF(Checklist!G4:G{n},"Pasó")/(COUNTA(Checklist!A4:A{n})-COUNTIF(Checklist!G4:G{n},"No aplica")),0)'])
res["B8"].number_format = "0%"; res.column_dimensions["A"].width = 18
lee = wb.create_sheet("Cómo usar")
for l in ["1. Abre https://pms-hoteles-atm.streamlit.app/ y pide a ATm Consulting las contraseñas temporales de prueba (se entregan aparte). Cada usuario crea su contraseña propia al primer ingreso.",
          "2. Haz cada escenario siguiendo la guía de pruebas (Word/PDF) o la columna Pasos.",
          "3. En la columna Resultado elige Pasó, Falló o No aplica.",
          "4. Si falló, escribe en Observaciones qué viste. Anota fecha y tu nombre.",
          "5. La hoja Resumen cuenta los resultados sola.",
          "Recuerda: todos los datos son ficticios; correos, pagos e integraciones son simulados."]:
    lee.append([l])
lee.column_dimensions["A"].width = 100
wb.save(DOCS / "checklist-pruebas-pms.xlsx")
print("xlsx ok", len(E), "escenarios")
