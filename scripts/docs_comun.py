"""Utilidades para generar los documentos Word/PDF del piloto (diagramas simples + estilos). Español sencillo."""
import re
import subprocess
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = Path(__file__).resolve().parent.parent
DOCS = RAIZ / "docs"
IMG = DOCS / "img_docs"
IMG.mkdir(exist_ok=True)
CAPT = DOCS / "capturas"
LOGO = RAIZ / "app" / "assets" / "atm_logo.png"
MARCA = "Prototipo diseñado para ATm Consulting"
AZUL = RGBColor(0x1A, 0x73, 0xE8)
GRIS = RGBColor(0x55, 0x55, 0x55)


# ---------- diagramas ----------
def diagrama(nombre, pasos, colores=None, ancho=None, alto=1.5, cols=None):
    """Flujo horizontal de cajas con flechas. pasos = lista de textos (usar \\n para partir líneas)."""
    n = len(pasos)
    cols = cols or n
    filas = -(-n // cols)
    w = ancho or min(2.4 * cols, 12)
    fig, ax = plt.subplots(figsize=(w, alto * filas + 0.2))
    ax.set_xlim(0, cols); ax.set_ylim(0, filas); ax.axis("off")
    pos = []
    for i, t in enumerate(pasos):
        f, c = divmod(i, cols)
        x, y = c + 0.5, filas - f - 0.5
        pos.append((x, y))
        col = (colores or {}).get(i, "#e8f0fe")
        ax.add_patch(FancyBboxPatch((x - 0.42, y - 0.3), 0.84, 0.6, boxstyle="round,pad=0.02,rounding_size=0.08",
                                    fc=col, ec="#1a73e8", lw=1.4))
        ax.text(x, y, t, ha="center", va="center", fontsize=9.5, color="#1f2937", linespacing=1.15)
    for i in range(n - 1):
        (x1, y1), (x2, y2) = pos[i], pos[i + 1]
        if y1 == y2:
            a = FancyArrowPatch((x1 + 0.44, y1), (x2 - 0.44, y2), arrowstyle="-|>", mutation_scale=14, color="#1a73e8", lw=1.6)
        else:
            a = FancyArrowPatch((x1, y1 - 0.32), (x2, y2 + 0.32), arrowstyle="-|>", mutation_scale=14, color="#1a73e8", lw=1.6,
                                connectionstyle="arc3,rad=0")
        ax.add_patch(a)
    fig.patch.set_facecolor("white")
    ruta = IMG / f"{nombre}.png"
    fig.savefig(ruta, dpi=170, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return ruta


# ---------- Word ----------
_EMOJI = re.compile("[\U0001F000-\U0001FFFF\u2600-\u27BF\uFE0F\u200d]+")


def lim(t):
    """Quita los emojis (no se ven bien en el PDF) y los espacios sobrantes."""
    return re.sub(r"\s{2,}", " ", _EMOJI.sub("", str(t))).replace("« ", "«").strip()

def nuevo_documento(titulo, subtitulo):
    d = Document()
    s = d.sections[0]
    s.left_margin = s.right_margin = Cm(2.2); s.top_margin = Cm(2); s.bottom_margin = Cm(2)
    st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    for nom, tam in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
        h = d.styles[nom]; h.font.name = "Calibri"; h.font.size = Pt(tam); h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F); h.font.bold = True
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    # pie de página
    p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(MARCA + " · datos 100 % ficticios"); r.font.size = Pt(8.5); r.font.color.rgb = GRIS
    # portada
    if LOGO.exists():
        d.add_picture(str(LOGO), width=Cm(4)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.LEFT
    for _ in range(3):
        d.add_paragraph()
    t = d.add_paragraph(); r = t.add_run(titulo); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    t = d.add_paragraph(); r = t.add_run(subtitulo); r.font.size = Pt(14); r.font.color.rgb = GRIS
    t = d.add_paragraph(); r = t.add_run(MARCA); r.italic = True; r.font.color.rgb = AZUL
    return d


def salto(d):
    d.add_page_break()


def parrafo(d, texto, negrita=False, cursiva=False, color=None, tam=None):
    p = d.add_paragraph(); r = p.add_run(lim(texto)); r.bold = negrita; r.italic = cursiva
    if color: r.font.color.rgb = color
    if tam: r.font.size = Pt(tam)
    p.paragraph_format.space_after = Pt(6)
    return p


def lista(d, items, numerada=False):
    for n, it in enumerate(items, 1):
        p = d.add_paragraph(style=None if numerada else "List Bullet")
        if numerada:
            p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
            r = p.add_run(f"{n}. "); r.bold = True
        if isinstance(it, tuple):
            r = p.add_run(lim(it[0])); r.bold = True; p.add_run(lim(it[1]))
        else:
            p.add_run(lim(it))
        p.paragraph_format.space_after = Pt(3)


def _sombra(celda, hexa):
    tcPr = celda._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexa); tcPr.append(sh)


def tabla(d, cabecera, filas, anchos=None):
    t = d.add_table(rows=1, cols=len(cabecera)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(cabecera):
        celda = t.rows[0].cells[i]; celda.text = ""; r = celda.paragraphs[0].add_run(lim(c)); r.bold = True; r.font.size = Pt(10)
        _sombra(celda, "E8F0FE")
    for f in filas:
        cs = t.add_row().cells
        for i, v in enumerate(f):
            cs[i].text = ""; r = cs[i].paragraphs[0].add_run(lim(v)); r.font.size = Pt(10)
    if anchos:
        for fila in t.rows:
            for i, a in enumerate(anchos):
                fila.cells[i].width = Cm(a)
    d.add_paragraph()
    return t


def imagen(d, ruta, ancho_cm=15.5, pie=None):
    d.add_picture(str(ruta), width=Cm(ancho_cm)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if pie:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(lim(pie)); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GRIS


def nota(d, texto, titulo="Ojo"):
    t = d.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; _sombra(c, "FFF8E1"); c.text = ""
    r = c.paragraphs[0].add_run(titulo + ": "); r.bold = True
    c.paragraphs[0].add_run(lim(texto))
    d.add_paragraph()


def guardar_y_pdf(d, nombre):
    ruta = DOCS / f"{nombre}.docx"
    d.save(ruta)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(DOCS), str(ruta)],
                   check=True, capture_output=True, timeout=180)
    return ruta, DOCS / f"{nombre}.pdf"
